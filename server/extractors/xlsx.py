from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from . import DocumentProcessingError, ExtractedDocument

MAX_ROWS_PER_SHEET = 10_000


def extract(base_text: str, file_bytes: bytes | None) -> ExtractedDocument:
    if not file_bytes:
        return ExtractedDocument(text=base_text)

    warnings: list[str] = []
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentProcessingError(400, 'XLSX preprocessing failed') from exc

    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines = [f'[sheet] {sheet.title}']
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= MAX_ROWS_PER_SHEET:
                    warnings.append(
                        f'Sheet {sheet.title} exceeds {MAX_ROWS_PER_SHEET} rows; '
                        'remaining rows were skipped.'
                    )
                    break
                cells = ['' if cell is None else str(cell) for cell in row]
                if any(cells):
                    lines.append('\t'.join(cells).rstrip())
            if len(lines) > 1:
                parts.append('\n'.join(lines))
    finally:
        workbook.close()

    text = '\n\n'.join(part for part in [base_text, *parts] if part)
    return ExtractedDocument(text=text, warnings=warnings)
