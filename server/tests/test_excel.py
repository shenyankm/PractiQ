"""Worksheet multimodal extraction, source association and durable retry."""

import hashlib
import json
from io import BytesIO
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from langgraph.checkpoint.memory import InMemorySaver
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from pydantic import ValidationError

from practiq_ai import config, execution, llm
from practiq_ai.contracts import DOCUMENT_MEDIA_TYPES, document_source_key
from practiq_ai.errors import DocumentProcessingError
from practiq_ai.extractors import xlsx
from practiq_ai.graphs import document, excel
from tests.support import (
    FakeModel,
    FakeObjectStore,
    local_graph,
    make_blank_pdf,
    make_image,
    question,
    run_config,
)


def workbook(*, images=False, charts=False, second=False, formula=False, empty=False):
    book = Workbook()
    sheet = book.worksheets[0]
    sheet.title = 'Quiz'
    if not empty:
        sheet.append(['1. What is shown?', 'Answer: a circle'])
        sheet.merge_cells('A2:C2')
        sheet['A2'] = 'Shared material'
    if images:
        sheet.add_image(ExcelImage(BytesIO(make_image())), 'D4')
    if formula:
        sheet['A3'] = '=1+1'
    if charts:
        sheet.append([2, 3])
        chart = BarChart()
        chart.title = 'Native Chart'
        chart.add_data(data=Reference(worksheet=sheet, min_col=1, max_col=2, min_row=3, max_row=3))
        sheet.add_chart(chart, 'F6')
        sheet.print_area = 'A1:B1'
    if second:
        other = book.create_sheet('Hidden')
        other['A1'] = 'Other sheet question'
        other.sheet_state = 'hidden'
        if charts:
            other.add_chart(BarChart(), 'D4')
    data = BytesIO()
    book.save(data)
    return data.getvalue()


def rewrite(data, transform):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, 'w') as dest:
        for item in source.infolist():
            dest.writestr(item, transform(item.filename, source.read(item)))
    return output.getvalue()


def setup(monkeypatch, payload, responses):
    digest = hashlib.sha256(payload).hexdigest()
    key = document_source_key('xlsx', digest)
    reference = {'sourceType': 'xlsx', 'objectKey': key, 'sha256': digest, 'sizeBytes': len(payload), 'mediaType': next(iter(DOCUMENT_MEDIA_TYPES['xlsx']))}
    files = FakeObjectStore({key: payload})
    model = FakeModel(responses=responses)
    kinds = []
    def get_model(kind='vision'):
        kinds.append(kind)
        assert kind == 'vision'
        return model
    monkeypatch.setattr(document, 'get_model', get_model)
    monkeypatch.setattr(document, 'get_object_store', lambda: files)
    return reference, model, files, kinds


def response(messages, schema):
    metadata = json.loads(messages[1].content)
    name = metadata['sheetName']
    assets = metadata['images']
    source = {'sheetName': name, 'cellRange': 'A1'}
    visuals = []
    if assets:
        source = {'sheetName': name, 'objectId': assets[0]['objectId']}
        visuals = [{'kind': 'image', 'description': 'Visible diagram', 'source': source, 'questionIndexes': [0]}]
    return {'questions': [{**question('Repeated'), 'excelSource': source}], 'visuals': visuals,
            'groups': [{'title': name, 'questionIndexes': [0]}]}


def test_extract_cells_images_anchors_and_hidden_formula():
    result = xlsx.extract(workbook(images=True, second=True, formula=True))
    first, second = result.worksheets
    assert '[merged] A2:C2' in first['text']
    assert 'A1=1. What is shown?' in first['text']
    assert first['assets'][0]['cellRange'] == 'D4'
    assert first['assets'][0]['original'] == make_image()
    assert first['assets'][0]['data'].startswith(b'\x89PNG')
    assert '[visibility] hidden' in second['text']
    assert result.truncated and 'no cached value' in result.warnings[0]


async def test_joint_sheet_parse_keeps_duplicates_and_remaps_associations(monkeypatch):
    reference, model, files, kinds = setup(monkeypatch, workbook(images=True, second=True), [response, response])
    graph = local_graph(InMemorySaver())
    settings = run_config()
    result = await graph.ainvoke({'document': reference}, settings)
    assert result['status'] == 'SUCCEEDED'
    assert len(result['result']['questions']) == 2
    assert [group['questionIndexes'] for group in result['result']['groups']] == [[0], [1]]
    assert result['result']['visualElements'][0]['questionIndexes'] == [0]
    assert result['result']['visualElements'][0]['page'] is None
    assert [s['excelSource']['sheetName'] for s in result['processing']['questionSources']] == ['Quiz', 'Hidden']
    assert not any(issue['code'] == 'SOURCE_TEXT_NOT_FOUND' for issue in result['processing']['quality']['issues'])
    assert len(model.calls) == len(result['usage']) == 2
    assert set(kinds) == {'vision'}
    checkpoint = await graph.aget_state(settings)
    manifests = [json.loads(files.blobs[ref['objectKey']]) for ref in checkpoint.values['chunkRefs']]
    assert manifests[0]['assets'][0]['original']['sizeBytes'] == len(make_image())
    assert len([m for m in next(call for call in model.calls if json.loads(call[1].content)['sheetName'] == 'Quiz') if isinstance(m.content, list)]) == 1
    await graph.ainvoke(None, settings)
    assert len(model.calls) == 2


async def test_image_only_workbook_can_produce_questions(monkeypatch):
    reference, model, _, _ = setup(monkeypatch, workbook(images=True, empty=True), [response])
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'SUCCEEDED'
    assert result['processing']['questionSources'][0]['excelSource']['objectId']
    assert len(model.calls) == 1


async def test_sheet_failure_retry_does_not_repeat_success(monkeypatch):
    reference, model, _, _ = setup(monkeypatch, workbook(second=True), [response] * 3)
    actual = excel.structured_call
    failed = True
    async def call(model, messages, schema, kind, **kwargs):
        if failed and json.loads(messages[1].content)['sheetName'] == 'Hidden':
            return None, [], 'OUTPUT_INVALID'
        return await actual(model, messages, schema, kind, **kwargs)
    monkeypatch.setattr(excel, 'structured_call', call)
    graph = local_graph(InMemorySaver())
    settings = run_config()
    first = await graph.ainvoke({'document': reference}, settings)
    assert first['status'] == 'PARTIAL'
    assert first['processing']['failures'][0]['index'] == 1
    failed = False
    result = await graph.ainvoke({'document': reference, 'retry': {'requestId': str(uuid4()), 'units': []}}, settings)
    assert result['status'] == 'SUCCEEDED'
    assert len(model.calls) == len(result['usage']) == 2
    assert [s['excelSource']['sheetName'] for s in result['processing']['questionSources']] == ['Quiz', 'Hidden']


@pytest.mark.parametrize('defect', ['wrong_sheet', 'unknown_object', 'bad_index', 'no_source', 'missing_source', 'visual_no_object'])
async def test_invalid_model_references_are_corrected_with_usage(monkeypatch, defect):
    def bad(messages, schema):
        result = response(messages, schema)
        source = result['questions'][0]['excelSource']
        if defect == 'wrong_sheet':
            source['sheetName'] = 'invented'
        elif defect == 'unknown_object':
            source['objectId'] = 'invented'
        elif defect == 'bad_index':
            result['visuals'][0]['questionIndexes'] = [99]
        elif defect == 'no_source':
            source.pop('objectId')
        elif defect == 'missing_source':
            result['questions'][0].pop('excelSource')
        else:
            source.pop('objectId')
            source['cellRange'] = 'A1'
        return result
    reference, model, _, _ = setup(monkeypatch, workbook(images=True), [bad, response])
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'SUCCEEDED'
    assert len(result['usage']) == len(model.calls) == 2


@pytest.mark.parametrize('mode,code,warning', [
    ('external', None, 'external or missing image was not loaded'),
    ('corrupt', None, 'XLSX_IMAGE_INVALID'),
    ('pixel_limit', 'XLSX_IMAGE_TOO_LARGE', 'Worksheet image exceeds pixel limit'),
    ('bytes_limit', 'DOCUMENT_PROCESSING_FAILED', 'Document visual content exceeds the configured limit'),
    ('row_limit', 'XLSX_SHEET_TOO_LARGE', 'Worksheet exceeds row limit'),
    ('text_limit', 'XLSX_SHEET_TOO_LARGE', 'Worksheet exceeds model input limit'),
    ('unsupported', None, 'Unsupported worksheet object: oleObjects'),
])
def test_visual_and_sheet_limits_are_reported(monkeypatch, mode, code, warning):
    data = workbook(images=True)
    if mode == 'external':
        data = rewrite(data, lambda path, value: value.replace(b'Target="/xl/media/image1.png"', b'Target="https://invalid.example/image.png" TargetMode="External"') if path.endswith('drawing1.xml.rels') else value)
    elif mode == 'corrupt':
        data = rewrite(data, lambda path, value: b'bad image' if path.endswith('image1.png') else value)
    elif mode == 'pixel_limit':
        monkeypatch.setenv('AI_MAX_VISION_PAGE_PIXELS', '1')
    elif mode == 'bytes_limit':
        monkeypatch.setenv('AI_MAX_VISION_BYTES', '1')
    elif mode == 'row_limit':
        monkeypatch.setattr(xlsx, 'MAX_ROWS_PER_SHEET', 1)
    elif mode == 'text_limit':
        monkeypatch.setenv('AI_MODEL_MAX_INPUT_CHARS', '100')
    else:
        data = rewrite(data, lambda path, value: value.replace(b'</worksheet>', b'<oleObjects/></worksheet>') if path.endswith('sheet1.xml') else value)
    result = xlsx.extract(data)
    assert result.truncated and result.warnings
    unit = result.worksheets[0]
    assert unit['failureCode'] == code
    assert any(warning in value for value in unit['warnings'])
    assert len(unit['assets']) == (1 if mode == 'unsupported' else 0)


@pytest.mark.parametrize('code', ['XLSX_CONVERTER_MISSING', 'XLSX_CONVERSION_FAILED', 'XLSX_CONVERSION_TIMEOUT'])
async def test_render_failure_preserves_other_sheet(monkeypatch, code):
    payload = workbook(charts=True, second=True)
    calls = 0
    def render(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise DocumentProcessingError(503, 'Renderer failed', code)
        return make_blank_pdf(1)
    monkeypatch.setattr(xlsx, 'convert_to_pdf', render)
    reference, model, _, _ = setup(monkeypatch, payload, [response])
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'PARTIAL'
    assert result['processing']['failures'][0]['code'] == code
    assert len(model.calls) == 1


def test_render_copy_keeps_drawings_and_clears_print_area():
    data = workbook(charts=True, second=True)
    with ZipFile(BytesIO(data)) as archive:
        copied = xlsx._render_copy(archive, 'xl/worksheets/sheet2.xml', 1)
        with ZipFile(BytesIO(copied)) as rendered:
            root = ET.fromstring(rendered.read('xl/workbook.xml'))
            states = [s.get('state') for s in root.findall(f'{{{xlsx.S}}}sheets/{{{xlsx.S}}}sheet')]
            assert states == ['hidden', 'visible']
            assert b'_xlnm.Print_Area' not in rendered.read('xl/workbook.xml')
            assert rendered.read('xl/charts/chart1.xml') == archive.read('xl/charts/chart1.xml')


@pytest.mark.parametrize('key', ['LLM_TEXT_MODEL', 'LLM_VISION_MODEL'])
def test_both_models_are_required(monkeypatch, key):
    monkeypatch.delenv(key)
    with pytest.raises(ValueError, match=key):
        config.load()


def test_models_and_execution_signatures_are_distinct(monkeypatch):
    llm.get_model.cache_clear()
    monkeypatch.setenv('LLM_TEXT_MODEL', 'text-only')
    monkeypatch.setenv('LLM_VISION_MODEL', 'image-capable')
    assert llm.get_model('text').model_name == 'text-only'
    assert llm.get_model('vision').model_name == 'image-capable'
    previous = execution.new_execution()
    monkeypatch.setenv('LLM_TEXT_MODEL', 'changed')
    with pytest.raises(DocumentProcessingError, match='original deployment'):
        execution.validate_execution(previous)
    with pytest.raises(ValueError, match='Unknown model'):
        llm.get_model('invalid')
    llm.get_model.cache_clear()


def test_visual_result_rejects_invalid_final_question_index():
    from practiq_ai.contracts import DocumentParseResult
    with pytest.raises(ValidationError, match='visual questionIndexes'):
        DocumentParseResult.model_validate({'questions': [question('Q')], 'groups': [], 'warnings': [], 'confidenceScore': 80,
                                          'visualElements': [{'kind': 'image', 'description': 'x', 'questionIndexes': [1]}]})


def with_shape(data):
    shape = f'''<xdr:twoCellAnchor xmlns:xdr="{xlsx.D}" xmlns:a="{xlsx.A}">
    <xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>20</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
    <xdr:to><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>25</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
    <xdr:sp><xdr:nvSpPr><xdr:cNvPr id="80" name="Native Shape"/><xdr:cNvSpPr/></xdr:nvSpPr>
    <xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="CC11AA"/></a:solidFill></xdr:spPr>
    <xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="1600"/><a:t>Native Shape</a:t></a:r></a:p></xdr:txBody>
    </xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'''
    def add(path, value):
        if path == 'xl/drawings/drawing1.xml':
            root = ET.fromstring(value)
            root.append(ET.fromstring(shape))
            return ET.tostring(root)
        return value
    return rewrite(data, add)


def test_real_calc_preserves_images_chart_shape_and_sheet_isolation(monkeypatch, libreoffice):
    import pypdfium2 as pdfium

    convert = xlsx.convert_to_pdf
    texts, image_counts = [], []
    def capture(data, **kwargs):
        pdf_bytes = convert(data, **kwargs)
        with pdfium.PdfDocument(pdf_bytes) as pdf:
            text = ''
            image_count = 0
            for page in pdf:
                textpage = page.get_textpage()
                try:
                    text += textpage.get_text_range()
                    image_count += sum(isinstance(obj, pdfium.PdfImage) for obj in page.get_objects())
                finally:
                    textpage.close()
                    page.close()
            texts.append(text)
            image_counts.append(image_count)
        return pdf_bytes
    monkeypatch.setattr(xlsx, 'convert_to_pdf', capture)
    result = xlsx.extract(with_shape(workbook(images=True, charts=True, second=True)))
    assert all(unit['failureCode'] is None for unit in result.worksheets), result.warnings
    # The source print area is A1:B1; chart F6 and shape B21:E26 must still render.
    assert 'Native Chart' in texts[0] and 'Native Shape' in texts[0]
    pages = [asset for asset in result.worksheets[0]['assets'] if asset['objectId'].startswith('render:')]
    assert pages
    assert any(Image.open(BytesIO(asset['data'])).convert('RGB').getextrema() != ((255, 255),) * 3 for asset in pages)
    assert 'Other sheet question' not in texts[0] and 'Other sheet question' in texts[1]
    assert 'Native Shape' not in texts[1]
    assert image_counts[0] >= 1
    assert result.worksheets[0]['objects'][-1]['cellRange'] == 'B21:E26'


@pytest.mark.parametrize('kind', ['text', 'csv'])
async def test_text_formats_route_to_text_model(monkeypatch, kind):
    from tests.support import source
    files, reference = source('1. Question?')
    if kind == 'csv':
        original = reference['objectKey']
        reference.update(sourceType='csv', mediaType='text/csv', objectKey=document_source_key('csv', reference['sha256']))
        files.blobs[reference['objectKey']] = files.blobs.pop(original)
    seen = []
    model = FakeModel(responses=[{'questions': [question('Question?')]}])
    def get_model(role='vision'):
        seen.append(role)
        assert role == 'text'
        return model
    monkeypatch.setattr(document, 'get_model', get_model)
    monkeypatch.setattr(document, 'get_object_store', lambda: files)
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'SUCCEEDED' and seen == ['text', 'text']


@pytest.mark.parametrize('mode', ['traversal', 'macro', 'doctype', 'bad_relationship', 'bad_anchor'])
def test_reject_unsafe_workbook_parts(mode):
    data = workbook(images=True)
    if mode == 'macro':
        out = BytesIO(data)
        with ZipFile(out, 'a') as archive:
            archive.writestr('xl/vbaProject.bin', b'unsafe')
        data = out.getvalue()
    elif mode == 'traversal':
        out = BytesIO(data)
        with ZipFile(out, 'a') as archive:
            archive.writestr('../escape', b'unsafe')
        data = out.getvalue()
    elif mode == 'doctype':
        data = rewrite(data, lambda path, value: b'<!DOCTYPE x>' + value if path.endswith('drawing1.xml') else value)
    elif mode == 'bad_relationship':
        data = rewrite(data, lambda path, value: value.replace(b'/xl/media/image1.png', b'../../../escape') if path.endswith('drawing1.xml.rels') else value)
    else:
        data = rewrite(data, lambda path, value: value.replace(b'<col>3</col>', b'<col>-1</col>') if path.endswith('drawing1.xml') else value)
    if mode in {'macro', 'traversal'}:
        with pytest.raises(DocumentProcessingError):
            xlsx.extract(data)
    else:
        assert xlsx.extract(data).worksheets[0]['failureCode']


@pytest.mark.parametrize('value', ['XFE1', 'A1048577', 'C3:A1'])
def test_excel_source_bounds(value):
    from practiq_ai.contracts import ExcelSource
    with pytest.raises(ValidationError, match='Invalid Excel'):
        ExcelSource(sheetName='Quiz', cellRange=value)


async def test_preparation_limit_is_a_nonretryable_sheet_failure(monkeypatch):
    payload = workbook(images=True, second=True)
    monkeypatch.setenv('AI_MAX_VISION_PAGE_PIXELS', '1')
    reference, model, _, _ = setup(monkeypatch, payload, [response])
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'PARTIAL'
    assert result['processing']['failures'][0]['code'] == 'XLSX_IMAGE_TOO_LARGE'
    assert result['processing']['failures'][0]['retryable'] is False
    assert len(model.calls) == 1


@pytest.mark.parametrize('code', ['AI_PROVIDER_ERROR', 'AI_USAGE_MISSING', 'MODEL_INPUT_TOO_LARGE'])
async def test_permanent_model_failure_keeps_successful_worksheet(monkeypatch, code):
    reference, model, _, _ = setup(monkeypatch, workbook(second=True), [response])
    actual = excel.structured_call
    async def call(model, messages, schema, kind, **kwargs):
        if json.loads(messages[1].content)['sheetName'] == 'Hidden':
            raise DocumentProcessingError(502, 'Failed', code)
        return await actual(model, messages, schema, kind, **kwargs)
    monkeypatch.setattr(excel, 'structured_call', call)
    result = await local_graph().ainvoke({'document': reference}, run_config())
    assert result['status'] == 'PARTIAL'
    assert result['processing']['failures'][0]['code'] == code
    assert len(result['usage']) == len(model.calls) == 1


async def test_execution_failure_is_not_swallowed_as_partial(monkeypatch):
    reference, _, _, _ = setup(monkeypatch, workbook(), [])
    async def fail(*args, **kwargs):
        raise DocumentProcessingError(503, 'Store unavailable', 'EXECUTION_STORE_UNAVAILABLE')
    monkeypatch.setattr(excel, 'structured_call', fail)
    with pytest.raises(DocumentProcessingError, match='Store unavailable'):
        await local_graph().ainvoke({'document': reference}, run_config())


def test_render_copy_disables_formulas_and_external_links():
    data = workbook(images=True, charts=True, formula=True)
    data = rewrite(data, lambda path, value: value.replace(b'Target="/xl/media/image1.png"', b'Target="https://invalid.example/x" TargetMode="External"') if path.endswith('drawing1.xml.rels') else value)
    with ZipFile(BytesIO(data)) as archive:
        copied = xlsx._render_copy(archive, 'xl/worksheets/sheet1.xml', 0)
    with ZipFile(BytesIO(copied)) as result:
        root = ET.fromstring(result.read('xl/worksheets/sheet1.xml'))
        assert root.find(f'.//{{{xlsx.S}}}f') is None
        assert b'https://invalid.example' not in result.read('xl/drawings/_rels/drawing1.xml.rels')


def test_transparent_images_are_composited_on_white():
    image = Image.new('RGBA', (2, 2), (0, 0, 0, 0))
    data = BytesIO()
    image.save(data, format='PNG')
    converted = Image.open(BytesIO(xlsx._png(data.getvalue())))
    assert converted.getpixel((0, 0)) == (255, 255, 255)


def test_sheet_preparation_deadline_bounds_multiple_conversions(monkeypatch):
    monkeypatch.setattr(xlsx, 'PREPARE_SECONDS', -1)
    monkeypatch.setattr(xlsx, 'convert_to_pdf', lambda *a, **k: pytest.fail('deadline must stop conversion'))
    result = xlsx.extract(workbook(charts=True, second=True))
    assert all(unit['failureCode'] == 'XLSX_PREPARE_TIMEOUT' for unit in result.worksheets)


def test_deadline_rechecked_after_building_render_copy(monkeypatch):
    clock = iter([0, 1, 151])
    monkeypatch.setattr(xlsx.time, 'monotonic', lambda: next(clock, 151))
    monkeypatch.setattr(xlsx, 'convert_to_pdf', lambda *a, **k: pytest.fail('deadline must stop conversion'))
    result = xlsx.extract(workbook(charts=True))
    assert result.worksheets[0]['failureCode'] == 'XLSX_PREPARE_TIMEOUT'
