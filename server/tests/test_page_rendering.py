"""Page ordering, conversion failures and real local rendering (no model calls)."""

import asyncio
import os
import shutil
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from practiq_ai.contracts import DOCUMENT_MEDIA_TYPES, document_source_key
from practiq_ai.errors import DocumentProcessingError
from practiq_ai.extractors import ExtractedDocument, docx, extract, pdf
from practiq_ai.graphs import document
from tests.test_documents import make_blank_pdf
from tests.test_workflows import FakeModel, local_graph, make_image, question, source


def paged_source(kind):
    store, reference = source("document")
    old_key = reference["objectKey"]
    reference.update(
        sourceType=kind,
        mediaType=next(iter(DOCUMENT_MEDIA_TYPES[kind])),
        objectKey=document_source_key(kind, reference["sha256"]),
    )
    store.blobs[reference["objectKey"]] = store.blobs.pop(old_key)
    return store, reference


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_missing_vision_fails_before_read_or_render(monkeypatch, kind):
    _, reference = paged_source(kind)
    monkeypatch.setattr(document, "get_model", lambda: None)
    monkeypatch.setattr(
        document, "get_object_store", lambda: pytest.fail("must not read")
    )
    monkeypatch.setattr(document, "extract", lambda *_: pytest.fail("must not render"))
    with pytest.raises(DocumentProcessingError) as error:
        asyncio.run(local_graph().ainvoke({"document": reference}))
    assert error.value.code == "VISION_MODEL_REQUIRED"


@pytest.mark.parametrize("kind", ["pdf", "docx"])
@pytest.mark.parametrize("failed", [set(), {1}, {0, 1, 2}])
def test_page_order_gaps_and_crops(monkeypatch, kind, failed):
    store, reference = paged_source(kind)
    model = FakeModel(responses=[{"questions": [question("Across pages")]}])
    monkeypatch.setattr(document, "get_object_store", lambda: store)
    monkeypatch.setattr(document, "get_model", lambda: model)
    monkeypatch.setattr(
        document,
        "extract",
        lambda *_: ExtractedDocument(text="", page_images=[make_image()] * 3),
    )

    async def page_call(_model, messages, schema, call_kind, *, runtime=None):
        prompt = messages[-1].content
        index = next(i for i in range(3) if f"PRIMARY page {i + 1}" in prompt)
        assert schema is document.PageParseResult and call_kind == "vision_parse"
        images = [message for message in messages if isinstance(message.content, list)]
        assert len(images) == (3 if index == 1 else 2)
        await asyncio.sleep((2 - index) * 0.01)
        if index in failed:
            return None, [], "OUTPUT_INVALID"
        return schema.model_validate({"questions": [question(f"Page {index + 1}")],
            "figures": [{"kind": "image", "description": "Figure", "bbox": [0.1, 0.1, 0.8, 0.8]}]}), [], None

    monkeypatch.setattr(document, "structured_call", page_call)
    if len(failed) == 3:
        with pytest.raises(DocumentProcessingError) as error:
            asyncio.run(local_graph().ainvoke({"document": reference}))
        assert error.value.code == "DOCUMENT_PARSE_FAILED"
        return
    result = asyncio.run(local_graph().ainvoke({"document": reference}))
    assert not model.calls  # No transcription or subsequent text-model call.
    assert [q["stem"] for q in result["result"]["questions"]] == [f"Page {i + 1}" for i in range(3) if i not in failed]
    assert result["status"] == ("PARTIAL" if failed else "SUCCEEDED")
    assert len(result["processing"]["failures"]) == len(failed)

    for figure in result["result"]["visualElements"]:
        assert figure["page"] not in failed
        assert figure["description"].startswith("[page crop]")
        assert figure["imageRef"]["objectKey"] in store.blobs


def converter_script(tmp_path, monkeypatch, body):
    executable = tmp_path / "converter"
    executable.write_text(f"#!{sys.executable}\n" + body)
    executable.chmod(0o700)
    monkeypatch.setenv("AI_SOFFICE_PATH", str(executable))
    return executable


def test_converter_missing(monkeypatch):
    monkeypatch.setenv("AI_SOFFICE_PATH", "/nonexistent/practiq-soffice")
    with pytest.raises(DocumentProcessingError) as error:
        docx.convert_to_pdf(b"test")
    assert error.value.code == "DOCX_CONVERTER_MISSING"


@pytest.mark.parametrize(
    "mode", ["success", "failure", "no_output", "timeout", "invalid_pdf"]
)
def test_converter_errors_and_cleanup(tmp_path, monkeypatch, mode):
    record = tmp_path / "directory"
    body = f"import sys, pathlib, time\np = pathlib.Path(sys.argv[-1]).parent\npathlib.Path({str(record)!r}).write_text(str(p))\n"
    if mode == "timeout":
        # A descendant must also be killed before TemporaryDirectory cleanup.
        body += "import subprocess\nchild = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])\n"
        body += f"pathlib.Path({str(tmp_path / 'child')!r}).write_text(str(child.pid))\ntime.sleep(60)\n"
        monkeypatch.setattr(docx, "CONVERSION_TIMEOUT_SECONDS", 3)
    elif mode == "failure":
        body += "sys.exit(1)\n"
    elif mode in {"success", "invalid_pdf"}:
        payload = make_blank_pdf(1) if mode == "success" else b"invalid PDF"
        body += f"(p / 'source.pdf').write_bytes({payload!r})\n"
    converter_script(tmp_path, monkeypatch, body)
    if mode in {"failure", "no_output", "timeout"}:
        with pytest.raises(DocumentProcessingError) as error:
            docx.convert_to_pdf(b"test")
        assert error.value.code == (
            "DOCX_CONVERSION_TIMEOUT" if mode == "timeout" else "DOCX_CONVERSION_FAILED"
        )
    else:
        result = docx.convert_to_pdf(b"test")
        if mode == "invalid_pdf":
            with pytest.raises(DocumentProcessingError):
                pdf.extract("", result)
        else:
            assert len(pdf.extract("", result).page_images) == 1
    assert not Path(record.read_text()).exists()
    if mode == "timeout":
        pid = int((tmp_path / "child").read_text())
        # An unreaped orphan may briefly be a zombie; it cannot still execute.
        import subprocess

        state = subprocess.run(
            ["ps", "-o", "stat=", "-p", str(pid)],
            capture_output=True,
            text=True,
            check=False,
        ).stdout.strip()
        assert not state or state.startswith("Z")


def test_converter_cannot_start(tmp_path, monkeypatch):
    executable = tmp_path / "broken"
    executable.write_bytes(b"not an executable")
    executable.chmod(0o700)
    monkeypatch.setenv("AI_SOFFICE_PATH", str(executable))
    with pytest.raises(DocumentProcessingError) as error:
        docx.convert_to_pdf(b"test")
    assert error.value.code == "DOCX_CONVERSION_FAILED"


def test_real_docx_rendering_and_original_image(monkeypatch, tmp_path):
    executable = shutil.which(os.environ.get("AI_SOFFICE_PATH") or "soffice")
    if executable is None:
        pytest.skip(
            "LibreOffice not installed; run this smoke check in the deployment environment"
        )
    monkeypatch.setenv("AI_SOFFICE_PATH", executable)
    original = docx.TemporaryDirectory
    paths = []

    def temporary_directory(**kwargs):
        directory = original(dir=tmp_path, **kwargs)
        paths.append(Path(directory.name))
        return directory

    monkeypatch.setattr(docx, "TemporaryDirectory", temporary_directory)
    image = BytesIO()
    Image.new("RGB", (40, 40), "red").save(image, "PNG")
    source = Document()
    source.add_paragraph("第一题：观察图形并回答问题。")
    source.add_picture(BytesIO(image.getvalue()))
    source.add_page_break()
    source.add_table(rows=1, cols=2).cell(0, 0).text = "第二页表格"
    payload = BytesIO()
    source.save(payload)
    result = extract("docx", payload.getvalue())
    assert result.text == ""
    assert len(result.page_images) == 2
    assert result.embedded_images == [image.getvalue()]
    assert paths and all(not path.exists() for path in paths)


@pytest.mark.parametrize("fixture", ["text-layer.pdf", "scanned.pdf"])
def test_real_pdf_fixtures_render_all_pages(fixture):
    import pypdfium2

    path = Path(__file__).parents[1] / "evals/fixtures/pdf" / fixture
    if not path.exists():
        pytest.fail(f"Missing PDF fixture: {path}")
    payload = path.read_bytes()
    with pypdfium2.PdfDocument(payload) as source:
        count = len(source)
    result = extract("pdf", payload)
    assert result.text == "" and len(result.page_images) == count


def test_mixed_pdf_renders_text_and_scanned_pages():
    import pypdfium2

    fixtures = Path(__file__).parents[1] / "evals/fixtures/pdf"
    with pypdfium2.PdfDocument.new() as merged:
        for name in ("text-layer.pdf", "scanned.pdf"):
            with pypdfium2.PdfDocument(fixtures / name) as source_pdf:
                merged.import_pages(source_pdf)
        count = len(merged)
        payload = BytesIO()
        merged.save(payload)
    result = extract("pdf", payload.getvalue())
    assert result.text == "" and len(result.page_images) == count


def test_docx_combined_visual_limit_and_bad_image_archive(monkeypatch):
    from zipfile import BadZipFile

    from tests.test_documents import make_docx

    payload = make_docx()
    monkeypatch.setattr(docx, "convert_to_pdf", lambda _: b"pdf")
    monkeypatch.setattr(
        docx,
        "extract_pdf",
        lambda *_: ExtractedDocument(text="", page_images=[b"123456"]),
    )
    monkeypatch.setattr(
        docx, "_extract_embedded_images", lambda *_: ([b"123456"], False)
    )
    monkeypatch.setenv("AI_MAX_VISION_BYTES", "10")
    with pytest.raises(DocumentProcessingError) as error:
        extract("docx", payload)
    assert error.value.status_code == 413

    def corrupt(*_args):
        raise BadZipFile("bad checksum")

    monkeypatch.setattr(docx, "_extract_embedded_images", corrupt)
    with pytest.raises(DocumentProcessingError, match="image extraction failed"):
        extract("docx", payload)
